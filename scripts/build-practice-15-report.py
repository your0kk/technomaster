"""Create the practice 15 report from recorded results and real screenshots."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'docs/practice-15'
OUT = ROOT / 'output/practice-15'
OUT.mkdir(parents=True, exist_ok=True)
facts = json.loads((DATA / 'results.json').read_text(encoding='utf-8'))
doc = Document()
section = doc.sections[0]
section.page_height = Mm(297)
section.page_width = Mm(210)
section.top_margin = section.bottom_margin = Mm(20)
section.left_margin = Mm(30)
section.right_margin = Mm(15)
section.footer_distance = Mm(10)
section.different_first_page_header_footer = True
for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(14 if name == 'Normal' else 16)
    style.paragraph_format.space_after = Pt(8)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.5
normal.first_line_indent = Mm(12.5)
normal.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def para(text):
    p = doc.add_paragraph(text)
    if 'https://' in text:
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = 0
    return p

def heading(text):
    doc.add_page_break()
    doc.add_heading(text, 1)

def code(text):
    for line in text.strip().splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = 0
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1
        p.paragraph_format.space_after = 0
        p.paragraph_format.keep_with_next = False
        r = p.add_run(line or ' ')
        r.font.name = 'Consolas'
        r.font.size = Pt(9)

def figure(name, caption, max_height=140):
    path = DATA / 'screenshots' / name
    with Image.open(path) as im:
        w, h = im.size
    width = min(160, max_height * w / h)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = 0
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Mm(width))
    p = doc.add_paragraph(caption)
    p.paragraph_format.first_line_indent = 0
    p.paragraph_format.line_spacing = 1
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(12)

def rows(headers, data, widths):
    label = 'Таблица 1 - Среда выполнения' if headers[0] == 'Среда' else 'Таблица 2 - Результаты проверки'
    p = doc.add_paragraph(label)
    p.paragraph_format.first_line_indent = 0
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.keep_with_next = True
    p.runs[0].font.size = Pt(12)
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for cells, values in [(t.rows[0].cells, headers)]:
        for c, v in zip(cells, values):
            c.text = v
    for values in data:
        for c, v in zip(t.add_row().cells, values):
            c.text = v
    for i, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit'); trpr.append(cant)
        if i == 0:
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
        for c, width in zip(row.cells, widths):
            c.width = Mm(width)
            pr = c._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ('top', 'left', 'bottom', 'right'):
                e = OxmlElement('w:' + edge)
                for k, v in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    e.set(qn('w:' + k), v)
                borders.append(e)
            pr.append(borders)
            if i == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'EAEAEA'); pr.append(shade)
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent = 0
                p.paragraph_format.line_spacing = 1.1
                p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(4)
                for r in p.runs:
                    r.font.size = Pt(11)
                    r.bold = i == 0

def centered(text, size=14, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = 0
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold

centered('Министерство образования Новосибирской области')
centered('ГБПОУ НСО Новосибирский авиационный технический колледж имени Б.С. Галущака')
doc.add_paragraph('\n\n')
p = doc.add_paragraph('ОТЧЕТ', 'Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
centered('по практическому занятию 15')
centered('Автоматизация сборки и деплоя проекта через GitHub Actions, GitLab и Docker', 16, True)
centered('на основе веб-приложения «ТехноМастер»')
doc.add_paragraph('\n\n')
p = doc.add_paragraph('Выполнил студент группы ПР-25.106\nУшаков Юрий Сергеевич')
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p.paragraph_format.first_line_indent = 0
doc.add_paragraph('\n\n\n')
centered('Новосибирск 2026')

heading('Цель и исходный проект')
para('Цель работы - собрать веб-приложение в Docker, автоматизировать проверку кода и публикацию образа через GitHub Actions. Для работы использован существующий проект «ТехноМастер»: сервис ремонта техники с каталогом, заказами и личными кабинетами.')
para('Приложение написано на Next.js и TypeScript. База данных и вход пользователей работают через Supabase. Проверки и контейнеризация добавлены к существующему репозиторию; страницы и пользовательские сценарии сохранены. Работа выполнена 6 октября 2026 года.')
rows(['Среда', 'Использованное значение'], [('Проект', 'your0kk/technomaster'), ('Node.js в сборке', '24'), ('Next.js', '16.3.6'), ('Docker Desktop', '4.94.0'), ('Docker Engine', '29.8.2, Linux amd64'), ('Runner Actions', 'ubuntu-latest'), ('Реестр образов', 'GitHub Container Registry')], [60, 100])
para('Практическая часть методички требует Docker и GitHub Actions. GitLab рассмотрен как аналог CI/CD: в нём конфигурация задаётся файлом .gitlab-ci.yml, а задания выполняет GitLab Runner. Отдельный конвейер GitLab в этой работе не создавался.')

heading('Схема автоматической сборки')
para('Изменения сначала отправлены в отдельную ветку. Pull Request запускает проверки кода и контейнера. При успешном результате изменения объединяются с main. Событие push в main дополнительно разрешает публикацию образа в GHCR.')
im = Image.new('RGB', (1400, 950), 'white'); draw = ImageDraw.Draw(im)
font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 29)
labels = ['Рабочая ветка и Pull Request', 'Проверки кода и тесты', 'Сборка и запуск контейнера', 'Успешные проверки и слияние в main', 'Push в main и публикация в GHCR', 'Pull образа и запуск на компьютере']
for i, text in enumerate(labels):
    y = 15 + i * 155
    draw.rectangle((100, y, 1300, y + 110), outline='#555555', width=3, fill='#F4F4F4')
    draw.text((700, y + 55), text, fill='black', font=font, anchor='mm')
    if i < 5:
        draw.line((700, y + 110, 700, y + 150), fill='black', width=3)
        draw.polygon([(692, y + 140), (708, y + 140), (700, y + 152)], fill='black')
im.save(DATA / 'screenshots/pipeline.png')
figure('pipeline.png', 'Рисунок 1 - Последовательность проверки и публикации', 150)
para('Публикуются два тега. latest указывает на последнюю сборку main. Полный хеш коммита позволяет выбрать конкретную версию независимо от последующих обновлений.')

heading('Dockerfile')
para('Dockerfile состоит из трёх стадий. Сначала устанавливаются зависимости через npm ci, затем выполняется production-сборка. В итоговый образ копируются standalone-сервер Next.js, статические файлы и public. Инструменты разработки в итоговом образе не нужны.')
code((ROOT / 'Dockerfile').read_text(encoding='utf-8'))
para('Параметр output со значением standalone добавлен в next.config.ts. Контейнер запускается командой node server.js, слушает порт 3000 и работает от пользователя node. HEALTHCHECK проверяет доступность страницы входа.')

heading('Контекст сборки и настройки окружения')
para('Файл .dockerignore исключает данные окружения, локальные пароли, кеш, Git, отчёты и файлы прототипа. Это уменьшает контекст сборки и не позволяет случайно включить закрытые файлы в Docker-образ. Ниже приведены основные правила; полный файл находится в репозитории.')
code('\n'.join((ROOT / '.dockerignore').read_text(encoding='utf-8').splitlines()[:19]))
para('NEXT_PUBLIC_SUPABASE_URL и NEXT_PUBLIC_SUPABASE_ANON_KEY нужны браузеру и задаются через build arguments. В GitHub они сохранены как repository variables. После изменения этих публичных настроек образ нужно пересобрать, потому что Next.js включает их в клиентский JavaScript.')
para('SUPABASE_SERVICE_ROLE_KEY не является аргументом сборки и не записывается в Dockerfile. При локальном запуске он передаётся через --env-file .env.local. Этот файл хранится только на компьютере. Для входа workflow в GHCR используется временный GITHUB_TOKEN с разрешением packages: write.')

heading('Локальная сборка и запуск')
para('В Docker Desktop запущен Linux Engine. Образ собран с именем technomaster:practice-15. Публичные настройки переданы сборщику из локального окружения, без серверного ключа. Затем контейнер запущен с файлом окружения и доступом только через локальный адрес.')
code('docker build --build-arg NEXT_PUBLIC_SUPABASE_URL \\\n  --build-arg NEXT_PUBLIC_SUPABASE_ANON_KEY \\\n  -t technomaster:practice-15 .\ndocker run -d --name technomaster-practice-15 \\\n  --env-file .env.local -p 127.0.0.1:3005:3000 \\\n  technomaster:practice-15\ndocker ps\ndocker stop technomaster-practice-15')
para('Команды приведены с переносами для чтения. Передача build arguments по имени предполагает, что публичные настройки уже заданы в окружении процесса. Для Windows команда может вводиться одной строкой. После проверки контейнер остановлен командой docker stop.')
figure('docker-desktop.png', 'Рисунок 2 - Работающий контейнер в Docker Desktop', 115)

heading('Проверка приложения в контейнере')
para('Скрипт scripts/check-container.mjs запускает временный контейнер, ждёт готовности сервера и проверяет восемь страниц. Главная, услуги, каталог, корзина, оформление заказа, вход и кабинеты сотрудников ответили HTTP 200. Анонимный запрос к административному API получил HTTP 401.')
para('Дополнительно проверены пользователь node, отсутствие серверного ключа в конфигурации образа и отсутствие .env, локальных файлов с реквизитами и приватных ключей внутри /app. В отдельном запуске с .env.local вход администратора и чтение заявок через контейнер дали HTTP 200. Данные базы при этой проверке не изменялись.')
figure('container-site.png', 'Рисунок 3 - Сайт из локального Docker-контейнера', 135)

heading('GitHub Actions и проверка кода')
para('Workflow находится в .github/workflows/ci.yml. Он запускается при push, pull_request и вручную. Задача check устанавливает Node.js, выполняет npm ci, линтер, проверку типов, пять тестов и production-сборку.')
workflow = (ROOT / '.github/workflows/ci.yml').read_text(encoding='utf-8')
code(workflow.split('  container:')[0])
para('Права по умолчанию ограничены чтением репозитория. needs: check не позволяет запустить сборку контейнера при ошибке проверки кода. Для рабочих веток устаревшие запуски отменяются, чтобы не тратить runner на предыдущий коммит.')

heading('Сборка и публикация образа')
para('Задача container использует Docker Buildx, собирает образ с публичными настройками и загружает его в локальный Docker runner. Перед публикацией запускается скрипт проверки контейнера. Кеш GitHub Actions ускоряет повторную сборку.')
code('  container:' + workflow.split('  container:')[1])
para('Вход в GHCR и docker push выполняются только для push в main. Pull Request проверяет образ без публикации. Благодаря этому пакет в реестре соответствует проверенному состоянию основной ветки.')

heading('Работа с Pull Request')
para('Для изменений создана ветка codex/practice-15-docker. После коммита и push открыт Pull Request №1 с целевой веткой main. В PR показаны изменённые файлы, описание контейнеризации и результаты проверок.')
code('git switch -c codex/practice-15-docker\ngit add Dockerfile .dockerignore next.config.ts \\\n  .github/workflows/ci.yml scripts/check-container.mjs\ngit commit -m "Добавить Docker и публикацию образа"\ngit push -u origin codex/practice-15-docker')
para('Обе задачи pull_request завершились успешно. После этого PR объединён с main. Такой порядок позволяет проверить изменения до их попадания в основную ветку и сохранить обсуждение, историю коммитов и результаты CI.')
figure('pull-request.png', 'Рисунок 4 - Pull Request №1 в репозитории', 115)
para('Ссылка: https://github.com/your0kk/technomaster/pull/1')

heading('Результат GitHub Actions')
para('Проверка PR и запуск после слияния завершились успешно. В запуске main выполнены обе задачи, включая авторизацию в GHCR и публикацию обоих тегов. Логи позволяют увидеть команду или шаг, на котором произошла бы ошибка.')
figure('actions.png', 'Рисунок 5 - Успешный workflow после слияния', 125)
para('Запуск PR: ' + facts['pr_run_url'])
para('Запуск main: ' + facts['main_run_url'])

heading('Образ в GitHub Container Registry')
para('Опубликован пакет ghcr.io/your0kk/technomaster. Для проверки образ получен из реестра, запущен на компьютере и проверен тем же скриптом. Проверка опубликованного образа завершилась успешно. Серверные настройки передаются при запуске и не входят в пакет.')
code('docker pull ghcr.io/your0kk/technomaster:latest\ndocker run -d --name technomaster-ghcr \\\n  --env-file .env.local -p 127.0.0.1:3006:3000 \\\n  ghcr.io/your0kk/technomaster:latest\ndocker ps\ndocker stop technomaster-ghcr')
figure('ghcr.png', 'Рисунок 6 - Пакет и версии образа в GHCR', 115)
para('Версия по коммиту: ' + facts['image_sha'])
para('Страница пакета: https://github.com/users/your0kk/packages/container/package/technomaster')

heading('Результаты и вывод')
rows(['Проверка', 'Результат'], [('Линтер и TypeScript', 'Ошибок нет'), ('Автоматические тесты', '5 из 5'), ('Production-сборка', 'Успешно'), ('Локальная сборка Docker', 'Успешно'), ('Страницы контейнера', '8 страниц, HTTP 200'), ('Healthcheck', 'healthy'), ('Анонимный административный API', 'HTTP 401'), ('Вход администратора в контейнере', 'HTTP 200'), ('Приватные файлы в образе', 'Не обнаружены'), ('Pull Request', '№1, проверки пройдены, слит'), ('Публикация GHCR', 'latest и полный хеш коммита'), ('Запуск образа из GHCR', 'Успешно')], [100, 60])
para('В ходе работы приложение упаковано в Docker и настроена автоматическая проверка перед публикацией образа. Образ можно получить из GHCR и запустить с локальными настройками окружения. Размещение сайта на сервере в эту работу не входило: публикация образа в реестре подготавливает его к такому запуску.')
para('Работа проверяет воспроизводимость сборки и основные HTTP-сценарии. Полное нагрузочное тестирование и аудит всех зависимостей не выполнялись. Предупреждения npm audit в существующих зависимостях требуют отдельного рассмотрения перед промышленным использованием.')

heading('Использованные источники')
for text in [
    'Практическое занятие №15 «Автоматизация сборки и деплоя проекта через GitHub Actions, GitLab и Docker». Методические рекомендации.',
    'Docker. Dockerfile reference. https://docs.docker.com/reference/dockerfile/',
    'Docker. Build context. https://docs.docker.com/build/building/context/',
    'GitHub. Publishing Docker images. https://docs.github.com/en/actions/tutorials/publish-packages/publish-docker-images',
    'GitHub. Creating a pull request. https://docs.github.com/en/pull-requests/collaborating-with-pull-requests/proposing-changes-to-your-work-with-pull-requests/creating-a-pull-request',
    'Next.js. Deploying. https://nextjs.org/docs/app/getting-started/deploying',
    'Исходный код проекта. https://github.com/your0kk/technomaster',
]:
    p = para(text); p.paragraph_format.first_line_indent = 0
para('Дата обращения к электронным источникам: 06.10.2026.')
doc.core_properties.author = 'Ушаков Юрий Сергеевич'
doc.core_properties.title = 'Практическое занятие 15 Контейнеризация ТехноМастер'
for style in doc.styles:
    if style.type == 1:
        style.font.name = 'Times New Roman'
        pr = style.element.find(qn('w:pPr'))
        if pr is not None:
            for border in list(pr.findall(qn('w:pBdr'))):
                pr.remove(border)
        fonts = style.element.find('.//' + qn('w:rFonts'))
        if fonts is not None:
            for key in list(fonts.attrib):
                if 'theme' in key.lower():
                    del fonts.attrib[key]
            for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
                fonts.set(qn('w:' + key), 'Times New Roman')
doc.save(OUT / 'technomaster-practice-15.docx')
print(OUT / 'technomaster-practice-15.docx')
